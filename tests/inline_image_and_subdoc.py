from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from docx import Document
from docxtpl import DocxTemplate, InlineImage


subdoc_path = Path("templates/merge_docx_subdoc.docx")
unsupported_image_subdoc_path = Path("templates/embedded_main_tpl.docx")

with TemporaryDirectory() as temporary_directory:
    output_dir = Path(temporary_directory)
    image_path = output_dir / "shared-image.png"
    template_path = output_dir / "inline_image_and_subdoc_tpl.docx"
    first_result = output_dir / "inline_image_and_subdoc_first.docx"
    second_result = output_dir / "inline_image_and_subdoc_second.docx"
    unsupported_image_template_path = output_dir / "unsupported_image_tpl.docx"
    unsupported_image_result = output_dir / "unsupported_image.docx"
    stream_result = output_dir / "inline_image_and_subdoc_stream.docx"

    with ZipFile(subdoc_path) as archive:
        image_path.write_bytes(archive.read("word/media/image1.png"))

    template_document = Document()
    template_document.add_paragraph("{{p subdoc }}")
    template_document.add_paragraph("{{ image }}")
    template_document.save(template_path)

    template = DocxTemplate(template_path)
    context = {
        "subdoc": template.new_subdoc(subdoc_path),
        "image": InlineImage(template, str(image_path)),
    }

    template.render(context)
    template.save(first_result)
    Document(first_result)

    template.render(context)
    template.save(second_result)
    Document(second_result)

    for result in (first_result, second_result):
        with ZipFile(result) as archive:
            media = [
                archive.read(name)
                for name in archive.namelist()
                if name.startswith("word/media/")
            ]
            assert sum(blob == image_path.read_bytes() for blob in media) == 1

    stream_template = DocxTemplate(template_path)
    image_stream = BytesIO(image_path.read_bytes())
    stream_template.render(
        {
            "subdoc": stream_template.new_subdoc(subdoc_path),
            "image": InlineImage(stream_template, image_stream),
        }
    )
    stream_template.save(stream_result)
    Document(stream_result)

    unsupported_image_template = Document()
    unsupported_image_template.add_paragraph("{{p subdoc }}")
    unsupported_image_template.save(unsupported_image_template_path)

    unsupported_image_document = DocxTemplate(unsupported_image_template_path)
    unsupported_image_document.render(
        {
            "subdoc": unsupported_image_document.new_subdoc(
                unsupported_image_subdoc_path
            )
        }
    )
    unsupported_image_document.save(unsupported_image_result)
    Document(unsupported_image_result)
