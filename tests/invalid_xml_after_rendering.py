from io import BytesIO

from docx import Document

from docxtpl import DocxTemplate


def create_template():
    document = Document()
    document.add_paragraph("Criticality: {{ defect.criticality }}")
    document.add_paragraph("Source: {{ defect.source }}")

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def render_template(autoescape=False):
    tpl = DocxTemplate(create_template())
    tpl.render(
        {"defect": {"criticality": "<select>", "source": "QA"}},
        autoescape=autoescape,
    )

    output = BytesIO()
    tpl.save(output)
    output.seek(0)
    return Document(output)


try:
    render_template()
except ValueError as exc:
    assert "Generated XML is invalid" in str(exc)
else:
    raise AssertionError("Unescaped XML should raise a rendering error")


document = render_template(autoescape=True)
assert [paragraph.text for paragraph in document.paragraphs] == [
    "Criticality: <select>",
    "Source: QA",
]
