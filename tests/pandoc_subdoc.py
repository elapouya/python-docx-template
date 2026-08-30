import io
import zipfile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

from docxtpl import DocxTemplate


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DEFAULT_NS = "urn:docxtpl:test-default"
ATTRIBUTE_NS = "urn:docxtpl:test-attribute"
CONFLICT_NS = "urn:docxtpl:test-conflict"


def make_template():
    document = Document()
    document.add_paragraph("{{p subdoc }}")
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    return DocxTemplate(stream)


def render_to_stream(template, subdoc):
    template.render({"subdoc": subdoc})
    stream = io.BytesIO()
    template.save(stream)
    stream.seek(0)
    return stream


def document_xml(stream):
    stream.seek(0)
    with zipfile.ZipFile(stream) as archive:
        return archive.read("word/document.xml")


template = make_template()
subdoc = template.new_subdoc("templates/pandoc_subdoc.docx")
output_stream = render_to_stream(template, subdoc)

document = Document(output_stream)
assert len(document.inline_shapes) == 1
image_relationships = [
    relationship
    for relationship in document.part.rels.values()
    if relationship.reltype == RT.IMAGE
]
assert len(image_relationships) == 1

root = etree.fromstring(document_xml(output_stream))
assert root.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}graphic") is not None
assert root.find(".//{http://schemas.openxmlformats.org/drawingml/2006/picture}pic") is not None
image_reference = root.find(
    ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
)
assert image_reference is not None
image_relationship_id = image_reference.get("{%s}embed" % R_NS)
assert image_relationship_id is not None
assert document.part.rels[image_relationship_id].reltype == RT.IMAGE

body = root.find("{%s}body" % W_NS)
section_properties = body.findall("{%s}sectPr" % W_NS)
assert len(section_properties) == 1
assert body[-1] is section_properties[0]


template = make_template()
subdoc = template.new_subdoc()
subdoc_body = subdoc.element.body
subdoc_body.insert(
    len(subdoc_body) - 1,
    etree.Element("{%s}node" % DEFAULT_NS, nsmap={None: DEFAULT_NS}),
)
attributed_paragraph = etree.Element(
    "{%s}p" % W_NS,
    nsmap={"attribute": ATTRIBUTE_NS},
)
attributed_paragraph.set("{%s}flag" % ATTRIBUTE_NS, "yes")
subdoc_body.insert(len(subdoc_body) - 1, attributed_paragraph)
subdoc_body.insert(
    len(subdoc_body) - 1,
    etree.Element("{%s}node" % CONFLICT_NS, nsmap={"w": CONFLICT_NS}),
)
output_stream = render_to_stream(template, subdoc)
root = etree.fromstring(document_xml(output_stream))

assert root.find(".//{%s}node" % DEFAULT_NS) is not None
assert any(
    element.get("{%s}flag" % ATTRIBUTE_NS) == "yes"
    for element in root.iter("{%s}p" % W_NS)
)
assert root.find(".//{%s}node" % CONFLICT_NS) is not None

body = root.find("{%s}body" % W_NS)
section_properties = body.findall("{%s}sectPr" % W_NS)
assert len(section_properties) == 1
assert body[-1] is section_properties[0]


template = make_template()
subdoc = template.new_subdoc()
for _ in range(10000):
    subdoc.add_paragraph("x")

fragment = subdoc._get_xml()
fragment_size = len(fragment.encode("utf-8"))
assert fragment_size < 1000000, (
    "subdocument fragment unexpectedly large: %d bytes" % fragment_size
)

fragment_body = etree.fromstring(
    ('<w:body xmlns:w="%s">%s</w:body>' % (W_NS, fragment)).encode("utf-8")
)
assert len(fragment_body) == 10000
assert all(element.tag == "{%s}p" % W_NS for element in fragment_body)
