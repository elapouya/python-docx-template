# -*- coding: utf-8 -*-

from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docxtpl import DocxTemplate

from _issue_621_ole import (
    W14_ANCHOR_ID,
    add_non_ole_vml_shape,
    add_ole_paragraph_anchor_ids,
    assert_file_identity,
    assert_ole_integrity,
)


SOURCE_FIXTURE = Path("templates/embedded_main_tpl.docx")
SOURCE_FIXTURE_SHA256 = (
    "40fa08a534110d411da00f38ed5ba971c2dcd12849e95edb3a3a9a48ca29919f"
)


assert_file_identity(SOURCE_FIXTURE, 164480, SOURCE_FIXTURE_SHA256)
with TemporaryDirectory() as temporary_directory:
    output_dir = Path(temporary_directory)
    anchored_source = output_dir / "anchored_source.docx"
    main_template = output_dir / "main_with_existing_ole.docx"
    result = output_dir / "embedded_ole_subdoc.docx"
    mixed_source = output_dir / "mixed_vml_and_ole.docx"
    mixed_main = output_dir / "mixed_vml_main.docx"
    mixed_result = output_dir / "mixed_vml_result.docx"

    source_anchor_count = add_ole_paragraph_anchor_ids(
        SOURCE_FIXTURE, anchored_source
    )
    assert source_anchor_count == 3
    main_document = Document(anchored_source)
    for index in range(source_anchor_count + 1, 11):
        paragraph = main_document.add_paragraph("reserved anchor %d" % index)
        anchor_id = "%08X" % index
        paragraph._p.set(
            W14_ANCHOR_ID,
            anchor_id.lower() if index == 10 else anchor_id,
        )
    main_document.add_paragraph("{{p first_subdoc }}")
    main_document.add_paragraph("{{p second_subdoc }}")
    main_document.save(main_template)

    template = DocxTemplate(main_template)
    template.render(
        {
            "first_subdoc": template.new_subdoc(anchored_source),
            "second_subdoc": template.new_subdoc(anchored_source),
        }
    )
    template.save(result)

    summary = assert_ole_integrity(
        result, expected_count=12, expected_anchor_count=16
    )
    assert summary["ole_count"] == 4 + (2 * 4)

    add_non_ole_vml_shape(
        anchored_source, mixed_source, "_x0000_i1025"
    )
    mixed_main_document = Document()
    mixed_main_document.add_paragraph("{{p mixed_subdoc }}")
    mixed_main_document.save(mixed_main)
    mixed_template = DocxTemplate(mixed_main)
    mixed_template.render(
        {"mixed_subdoc": mixed_template.new_subdoc(mixed_source)}
    )
    mixed_template.save(mixed_result)
    mixed_summary = assert_ole_integrity(
        mixed_result, expected_count=4, expected_anchor_count=3
    )
    assert mixed_summary["ole_count"] == 4
