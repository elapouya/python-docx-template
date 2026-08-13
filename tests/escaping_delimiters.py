from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate


output_dir = Path(__file__).parent / "output"
output_dir.mkdir(exist_ok=True)
template_path = output_dir / "escaping_delimiters_split_tpl.docx"
result_path = output_dir / "escaping_delimiters_split.docx"

escaped_delimiters = ("{_%", "%_}", "{_{", "}_}")
expected_delimiters = ("{%", "%}", "{{", "}}")

template = Document()

for delimiter in escaped_delimiters:
    template.add_paragraph(delimiter)

template.add_paragraph()

for delimiter in escaped_delimiters:
    paragraph = template.add_paragraph()
    for character in delimiter:
        paragraph.add_run(character)

# These are not escaped delimiters and must not be changed.
template.add_paragraph("{a_{hello")
template.add_paragraph("{f_{")
template.save(template_path)

doc = DocxTemplate(template_path)
doc.render({})
doc.save(result_path)

paragraphs = [paragraph.text for paragraph in Document(result_path).paragraphs]
assert paragraphs == [
    *expected_delimiters,
    "",
    *expected_delimiters,
    "{a_{hello",
    "{f_{",
]
